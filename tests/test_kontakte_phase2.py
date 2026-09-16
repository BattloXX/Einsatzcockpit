"""Integrationstests fuer das zentrale Kontakte-Modul Phase 2."""

from __future__ import annotations

import json
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook
from PIL import Image

from app.core.security import hash_password
from app.core.telefon import telefon_normalisiert
from app.core.tenant import set_tenant_context
from app.db import SessionLocal
from app.models.kontakt import Kontakt, KontaktKategorie, ObjektKontaktFreigabe
from app.models.master import FireDept, OrgSettings, SystemSettings
from app.models.objekt import Objekt, ObjektKontakt
from app.models.user import Role, User, UserRole
from app.services import kontakt_service
from app.services.kontakt_transfer_service import apply_preview, export_xlsx, parse_import, preview_import, save_preview


def _rolle(db, code: str) -> Role:
    rolle = db.query(Role).filter(Role.code == code).first()
    assert rolle is not None
    return rolle


def _setup_user(username: str, rolle: str, *, org_id: int = 1) -> User:
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        user = User(
            username=username,
            password_hash=hash_password("Test1234!"),
            display_name=username,
            org_id=org_id,
            active=True,
        )
        db.add(user)
        db.flush()
        db.add(UserRole(user_id=user.id, role_id=_rolle(db, rolle).id))
        system = db.get(SystemSettings, "kontakte_module_enabled")
        if system is None:
            db.add(SystemSettings(key="kontakte_module_enabled", value="true"))
        else:
            system.value = "true"
        settings = db.query(OrgSettings).filter_by(org_id=org_id).first()
        if settings is None:
            settings = OrgSettings(org_id=org_id)
            db.add(settings)
        settings.kontakte_module_enabled = True
        db.commit()
        db.refresh(user)
        db.expunge(user)
        return user
    finally:
        db.close()


def _login(client, username: str) -> None:
    client.cookies.clear()
    client.get("/login")
    client.post(
        "/login",
        data={"username": username, "password": "Test1234!", "_csrf": client.cookies.get("ec_csrf")},
        follow_redirects=False,
    )


def _post_data(**werte):
    return werte


def test_neu_dialog_ist_leer_und_verwendet_den_anlage_endpoint(client):
    user = _setup_user("kontakte_neu_dialog", "kontakt_verwalter")
    _login(client, user.username)

    response = client.get("/kontakte/neu")

    assert response.status_code == 200
    assert '<dialog id="kontaktModal" class="modal" open>' in response.text
    assert 'action="/kontakte/"' in response.text
    assert 'enctype="multipart/form-data"' in response.text
    assert 'name="profilbild"' in response.text
    assert "Kontakt bearbeiten" not in response.text


def test_bearbeiten_dialog_verwendet_update_endpoint_und_vorhandene_daten(client):
    user = _setup_user("kontakte_bearbeiten_dialog", "kontakt_verwalter")
    _login(client, user.username)
    csrf = client.cookies.get("ec_csrf")
    erstellt = client.post(
        "/kontakte/",
        data={
            "_csrf": csrf,
            "typ": "person",
            "anzeigename": "Vorhandener Kontakt",
            "vorname": "Vorhanden",
            "nummer": ["+43 664 111", "+43 664 222"],
            "telefon_label": ["Mobil", "Dienst"],
            "bevorzugt": ["0"],
            "sms_eignung": ["1"],
        },
        follow_redirects=False,
    )
    assert erstellt.status_code == 303
    kontakt_id = int(erstellt.headers["location"].rsplit("/", 1)[1])

    response = client.get(f"/kontakte/{kontakt_id}/bearbeiten")

    assert response.status_code == 200
    assert f'action="/kontakte/{kontakt_id}"' in response.text
    assert 'value="Vorhandener Kontakt"' in response.text
    assert 'value="Vorhanden"' in response.text
    assert "+43 664 111" in response.text
    assert "+43 664 222" in response.text
    assert "Mobil" in response.text and "Dienst" in response.text
    aktualisiert = client.post(
        f"/kontakte/{kontakt_id}",
        data={"_csrf": csrf, "version": "0", "typ": "person", "anzeigename": "Aktualisierter Kontakt"},
        follow_redirects=False,
    )
    assert aktualisiert.status_code == 303
    db = SessionLocal()
    set_tenant_context(db, user.org_id)
    try:
        assert db.query(Kontakt).filter_by(anzeigename="Aktualisierter Kontakt").count() == 1
        assert db.query(Kontakt).filter_by(anzeigename="Vorhandener Kontakt").count() == 0
    finally:
        db.close()


def test_bearbeiten_dialog_liefert_telefone_als_json_insel(client):
    user = _setup_user("kontakte_telefon_json_insel", "kontakt_verwalter")
    _login(client, user.username)
    erstellt = client.post(
        "/kontakte/",
        data={
            "_csrf": client.cookies.get("ec_csrf"),
            "typ": "person",
            "anzeigename": "JSON Telefon",
            "nummer": ["+43 664 1234567"],
            "telefon_label": ["Mobil"],
        },
        follow_redirects=False,
    )
    kontakt_id = int(erstellt.headers["location"].rsplit("/", 1)[1])

    response = client.get(f"/kontakte/{kontakt_id}/bearbeiten")

    assert response.status_code == 200
    match = re.search(
        r'<script type="application/json" id="kontakt-telefone-daten">(.*?)</script>', response.text
    )
    assert match is not None
    assert json.loads(match.group(1)) == [
        {"label": "Mobil", "nummer": "+43 664 1234567", "bevorzugt": False, "sms": False, "smsManuell": True}
    ]


def test_bearbeiten_redisplay_normalisiert_none_nach_versionskonflikt(client):
    user = _setup_user("kontakte_redisplay_none", "kontakt_verwalter")
    _login(client, user.username)
    erstellt = client.post(
        "/kontakte/",
        data={"_csrf": client.cookies.get("ec_csrf"), "typ": "person", "anzeigename": "Konflikt Kontakt"},
        follow_redirects=False,
    )
    kontakt_id = int(erstellt.headers["location"].rsplit("/", 1)[1])

    response = client.post(
        f"/kontakte/{kontakt_id}",
        data={
            "_csrf": client.cookies.get("ec_csrf"),
            "version": "999",
            "typ": "person",
            "anzeigename": "Konflikt Kontakt",
            "vorname": "None",
        },
    )

    assert response.status_code == 200
    assert 'value="None"' not in response.text


def test_detail_vollaufruf_oeffnet_keinen_dialog_und_rendert_none_nicht(client):
    user = _setup_user("kontakte_detail_reload", "kontakt_verwalter")
    db = SessionLocal()
    set_tenant_context(db, user.org_id)
    try:
        kontakt = kontakt_service.create_kontakt(
            db,
            {"typ": "person", "anzeigename": "Fischnaller"},
            [],
            [],
            org_id=user.org_id,
            user_id=None,
        )
        kontakt_id = kontakt.id
    finally:
        db.close()
    _login(client, user.username)

    response = client.get(f"/kontakte/{kontakt_id}")

    assert response.status_code == 200
    assert '<dialog id="kontaktModal" class="modal">' in response.text
    assert '<dialog id="kontaktModal" class="modal" open>' not in response.text
    assert 'value="None"' not in response.text
    assert '>None</textarea>' not in response.text


def test_kontakt_detail_trennt_objektrolle_optisch(client):
    user = _setup_user("kontakte_objektrolle_badge", "kontakt_verwalter")
    _login(client, user.username)
    erstellt = client.post(
        "/kontakte/",
        data={"_csrf": client.cookies.get("ec_csrf"), "typ": "person", "anzeigename": "Objekt Kontakt"},
        follow_redirects=False,
    )
    kontakt_id = int(erstellt.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        objekt = Objekt(org_id=user.org_id, nummer=999901, name="Holzbau Berchtold")
        db.add(objekt)
        db.flush()
        db.add(ObjektKontakt(org_id=user.org_id, objekt_id=objekt.id, kontakt_id=kontakt_id, art="sonstig"))
        db.commit()
        objekt_nummer = objekt.anzeige_nummer
    finally:
        db.close()

    response = client.get(f"/kontakte/{kontakt_id}")

    assert response.status_code == 200
    assert (
        f'<strong>{objekt_nummer} · Holzbau Berchtold</strong> '
        '<span class="badge-pill badge-pill--gray">Sonstig</span>'
    ) in response.text


def test_profilbild_wird_klein_gespeichert_und_ausgeliefert(client, tmp_path, monkeypatch):
    user = _setup_user("kontakte_profilbild", "kontakt_verwalter")
    _login(client, user.username)
    monkeypatch.setattr("app.routers.ui_kontakt._BILD_DIR", tmp_path)

    source = BytesIO()
    Image.new("RGB", (1600, 900), "red").save(source, "PNG")
    response = client.post(
        "/kontakte/",
        data={
            "_csrf": client.cookies.get("ec_csrf"),
            "typ": "person",
            "anzeigename": "Bild Kontakt",
        },
        files={"profilbild": ("portrait.png", source.getvalue(), "image/png")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    kontakt_id = int(response.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    set_tenant_context(db, 1)
    try:
        kontakt = kontakt_service.get_kontakt(db, kontakt_id)
        assert kontakt is not None and kontakt.bild_pfad
        datei = tmp_path / kontakt.bild_pfad
        assert datei.is_file()
        with Image.open(datei) as bild:
            assert bild.format == "JPEG"
            assert max(bild.size) <= 192
    finally:
        db.close()

    bild_response = client.get(f"/kontakte/{kontakt_id}/profilbild")
    assert bild_response.status_code == 200
    assert bild_response.headers["content-type"] == "image/jpeg"
    assert f'src="/kontakte/{kontakt_id}/profilbild"' in client.get("/kontakte").text
    assert f'src="/kontakte/{kontakt_id}/profilbild"' in client.get(f"/kontakte/{kontakt_id}").text


def test_crud_multitelefon_kategorien_und_konflikt(client):
    user = _setup_user("kontakte_crud", "kontakt_verwalter")
    _login(client, user.username)
    csrf = client.cookies.get("ec_csrf")
    payload = _post_data(
        _csrf=csrf,
        typ="person",
        anzeigename="Anna Beispiel",
        organisation="FF Beispiel",
        nummer=["00 43 664 123456", "+43 555 777"],
        telefon_label=["Mobil", "Dienst"],
        bevorzugt=["1"],
        sms_eignung=["0"],
        kategorien="Einsatzleitung, Atemschutz",
    )
    response = client.post("/kontakte/", data=payload, follow_redirects=False)
    assert response.status_code == 303
    kontakt_id = int(response.headers["location"].rsplit("/", 1)[1])

    db = SessionLocal()
    set_tenant_context(db, 1)
    try:
        kontakt = kontakt_service.get_kontakt(db, kontakt_id)
        assert kontakt is not None
        assert [telefon.nummer_normalisiert for telefon in kontakt.telefone] == ["+43664123456", "+43555777"]
        assert kontakt.telefone[1].bevorzugt and kontakt.telefone[0].sms_eignung
        assert {z.kategorie.name for z in kontakt.kategorien} == {"Einsatzleitung", "Atemschutz"}
        version = kontakt.version
    finally:
        db.close()

    update = _post_data(
        _csrf=csrf,
        version=str(version),
        typ="person",
        anzeigename="Anna Aktualisiert",
        nummer=["+43 555 777"],
        telefon_label=["Dienst"],
        bevorzugt=["0"],
        sms_eignung=["0"],
        kategorien="Einsatzleitung",
    )
    assert client.post(f"/kontakte/{kontakt_id}", data=update, follow_redirects=False).status_code == 303
    stale = _post_data(_csrf=csrf, version=str(version), typ="person", anzeigename="Stale")
    conflict = client.post(f"/kontakte/{kontakt_id}", data=stale)
    assert conflict.status_code == 200
    assert "inzwischen geaendert" in conflict.text
    assert "Stale" in conflict.text

    db = SessionLocal()
    set_tenant_context(db, 1)
    try:
        kontakt = kontakt_service.get_kontakt(db, kontakt_id)
        assert kontakt is not None and kontakt.anzeigename == "Anna Aktualisiert"
        assert len(kontakt.telefone) == 1 and kontakt.telefone[0].bevorzugt
        aktuelle_version = kontakt.version
    finally:
        db.close()
    assert (
        client.post(f"/kontakte/{kontakt_id}/archivieren", data={"_csrf": csrf}, follow_redirects=False).status_code
        == 303
    )
    assert "Anna Aktualisiert" not in client.get("/kontakte").text
    assert aktuelle_version > version


def test_rolle_suche_kategorie_idempotent_und_tenant_isolation(client):
    verwalter = _setup_user("kontakte_objekt", "objekt_verwalter")
    readonly = _setup_user("kontakte_readonly", "readonly")
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        org_b = FireDept(slug="kontakte-org-b", name="Kontakte Org B", color="#112233", bos="FW")
        db.add(org_b)
        db.flush()
        org_b_id = org_b.id
        db.commit()
    finally:
        db.close()
    _setup_user("kontakte_org_b", "kontakt_verwalter", org_id=org_b_id)

    _login(client, verwalter.username)
    csrf = client.cookies.get("ec_csrf")
    for name in ("Org A Kontakt", "Org A Zweitkontakt"):
        assert (
            client.post(
                "/kontakte/",
                data=_post_data(
                    _csrf=csrf,
                    typ="person",
                    anzeigename=name,
                    nummer=["00 43 664 123456"],
                    telefon_label=["Mobil"],
                    kategorien="Gemeinsam",
                    duplikate_bestaetigt="1" if name == "Org A Zweitkontakt" else "",
                ),
                follow_redirects=False,
            ).status_code
            == 303
        )
    db = SessionLocal()
    set_tenant_context(db, 1)
    try:
        assert db.query(KontaktKategorie).filter(KontaktKategorie.name == "Gemeinsam").count() == 1
        kontakt_id = db.query(Kontakt).filter(Kontakt.anzeigename == "Org A Kontakt").one().id
    finally:
        db.close()
    assert "Org A Kontakt" in client.get("/kontakte/liste?q=%2B43664123456").text

    _login(client, readonly.username)
    assert client.get("/kontakte").status_code == 200
    assert client.get("/kontakte/liste").status_code == 200
    assert client.get(f"/kontakte/{kontakt_id}").status_code == 200
    assert client.post("/kontakte/", data={"_csrf": client.cookies.get("ec_csrf")}).status_code == 403

    _login(client, "kontakte_org_b")
    assert client.get(f"/kontakte/{kontakt_id}").status_code == 404
    assert "Org A Kontakt" not in client.get("/kontakte/liste?q=Org%20A%20Kontakt").text


def test_kontakte_module_toggle_returns_404(client):
    user = _setup_user("kontakte_toggle", "kontakt_verwalter")
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        db.query(OrgSettings).filter_by(org_id=user.org_id).one().kontakte_module_enabled = False
        db.commit()
    finally:
        db.close()
    _login(client, user.username)
    assert client.get("/kontakte").status_code == 404
    assert client.get("/kontakte/liste").status_code == 404
    assert client.post("/kontakte/", data={"_csrf": client.cookies.get("ec_csrf")}).status_code == 404


def test_kontakt_export_ist_org_gebunden_und_enthaelt_alle_xlsx_blaetter(client):
    user = _setup_user("kontakte_export", "kontakt_verwalter")
    _login(client, user.username)
    csrf = client.cookies.get("ec_csrf")
    assert (
        client.post(
            "/kontakte/",
            data=_post_data(_csrf=csrf, typ="person", anzeigename="Export Kontakt", nummer=["+43 664 1"]),
            follow_redirects=False,
        ).status_code
        == 303
    )
    csv_response = client.get("/kontakte/export.csv")
    assert csv_response.status_code == 200
    assert "Export Kontakt" in csv_response.content.decode("utf-8-sig")
    xlsx_response = client.get("/kontakte/export.xlsx")
    assert xlsx_response.status_code == 200
    workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
    assert workbook.sheetnames == ["Kontakte", "Telefonnummern", "Objektzuordnungen", "Anleitung"]


def test_import_vorschau_uebernimmt_telefon_und_liefert_ergebnis_csv(client):
    user = _setup_user("kontakte_import", "kontakt_verwalter")
    _login(client, user.username)
    csrf = client.cookies.get("ec_csrf")
    csv_data = (
        "version;id;typ;anzeigename;organisation;funktion;email;erreichbarkeit\n"
        "1;;person;Import Kontakt;Import GmbH;Bereitschaft;import@example.test;tagsueber\n"
    )
    response = client.post(
        "/kontakte/import/vorschau",
        data={"_csrf": csrf},
        files={"datei": ("kontakte.csv", csv_data, "text/csv")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    preview_url = response.headers["location"]
    assert client.get(preview_url).status_code == 200
    assert client.post(f"{preview_url}/uebernehmen", data={"_csrf": csrf}, follow_redirects=False).status_code == 303
    result = client.get(f"{preview_url}/ergebnis.csv")
    assert result.status_code == 200
    assert "uebernommen" in result.content.decode("utf-8-sig")
    assert client.get("/kontakte/vorlage.xlsx").status_code == 200
    assert client.get("/kontakte/vorlage.csv?beispiel=1").status_code == 200


def test_xlsx_roundtrip_uebernimmt_telefone_und_objektzuordnungen(client):
    user = _setup_user("kontakte_xlsx_roundtrip", "kontakt_verwalter")
    _login(client, user.username)
    csrf = client.cookies.get("ec_csrf")
    assert (
        client.post(
            "/kontakte/",
            data=_post_data(_csrf=csrf, typ="person", anzeigename="Roundtrip Kontakt", nummer=["+43 664 111"]),
            follow_redirects=False,
        ).status_code
        == 303
    )
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        kontakt = db.query(Kontakt).filter_by(org_id=user.org_id, anzeigename="Roundtrip Kontakt").one()
        objekt = Objekt(org_id=user.org_id, nummer=981234, name="Roundtrip Objekt")
        db.add(objekt)
        db.flush()
        zuordnung = ObjektKontakt(org_id=user.org_id, objekt_id=objekt.id, kontakt_id=kontakt.id)
        db.add(zuordnung)
        db.flush()
        db.add(
            ObjektKontaktFreigabe(
                org_id=user.org_id,
                objekt_kontakt_id=zuordnung.id,
                kanal="sms",
                ziel_wert=telefon_normalisiert(kontakt.telefone[0].nummer),
                aktiv=True,
            )
        )
        db.commit()
        rows = parse_import(export_xlsx(db, user.org_id), "kontakte.xlsx")
        row = next(row for row in rows if row["id"] == str(kontakt.id))
        assert row["_telefone"][0]["nummer"] == kontakt.telefone[0].nummer
        assert row["_zuordnungen"][0]["objekt_id"] == str(objekt.id)
        row["funktion"] = "Aktualisiert"
        row["_telefone"][0]["label"] = "Mobil"
        row["_zuordnungen"][0]["rolle"] = "betreiber"
        preview = preview_import(db, user.org_id, rows)
        entry = save_preview(db, user.org_id, user.id, preview)
        assert apply_preview(db, user.org_id, user.id, entry.id) == 1
        db.expire_all()
        assert db.get(Kontakt, kontakt.id).telefone[0].label == "Mobil"
        assert db.query(ObjektKontakt).filter_by(objekt_id=objekt.id, kontakt_id=kontakt.id).one().art == "betreiber"
        assert db.query(ObjektKontaktFreigabe).filter_by(objekt_kontakt_id=zuordnung.id).count() == 1
    finally:
        db.close()


def test_xlsx_import_erkennt_mobilnummer_ohne_sms_spalte():
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        kontakt = kontakt_service.create_kontakt(
            db,
            {"typ": "person", "anzeigename": "XLSX Mobil", "funktion": "Alt"},
            [],
            [],
            org_id=1,
            user_id=None,
        )
        workbook = Workbook()
        kontakte = workbook.active
        kontakte.title = "Kontakte"
        kontakte.append(["id", "anzeigename", "typ", "funktion"])
        kontakte.append([str(kontakt.id), "XLSX Mobil", "person", "Neu"])
        telefone = workbook.create_sheet("Telefonnummern")
        telefone.append(["kontakt_id", "nummer", "label", "bevorzugt"])
        telefone.append([str(kontakt.id), "+43 664 123456", "Mobil", "0"])
        content = BytesIO()
        workbook.save(content)

        rows = parse_import(content.getvalue(), "kontakte.xlsx")
        assert "sms_eignung" not in rows[0]["_telefone"][0]
        preview = preview_import(db, 1, rows)
        entry = save_preview(db, 1, 1, preview)
        assert apply_preview(db, 1, 1, entry.id) == 1
        db.expire_all()
        assert db.get(Kontakt, kontakt.id).telefone[0].sms_eignung is True
    finally:
        db.close()


def test_nummernaenderung_wirkt_an_beiden_objekten_ohne_freigabe_wanderung():
    nummer_a = "+43 664 111 222"
    nummer_b = "+43 664 333 444"
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        kontakt = kontakt_service.create_kontakt(
            db,
            {"typ": "person", "anzeigename": "Gemeinsamer Kontakt"},
            [{"nummer": nummer_a}],
            [],
            org_id=1,
            user_id=None,
        )
        objekte = [
            Objekt(org_id=1, nummer=981235, name="Nummernaenderung Objekt 1"),
            Objekt(org_id=1, nummer=981236, name="Nummernaenderung Objekt 2"),
        ]
        db.add_all(objekte)
        db.flush()
        zuordnungen = [
            ObjektKontakt(org_id=1, objekt_id=objekt.id, kontakt_id=kontakt.id)
            for objekt in objekte
        ]
        db.add_all(zuordnungen)
        db.flush()
        db.add_all(
            ObjektKontaktFreigabe(
                org_id=1,
                objekt_kontakt_id=zuordnung.id,
                kanal="sms",
                ziel_wert=telefon_normalisiert(nummer_a),
                aktiv=True,
            )
            for zuordnung in zuordnungen
        )
        db.commit()

        kontakt_service.update_kontakt(
            db,
            kontakt.id,
            {"typ": "person", "anzeigename": "Gemeinsamer Kontakt"},
            [{"nummer": nummer_b}],
            [],
            version=kontakt.version,
            org_id=1,
            user_id=None,
        )
        db.expire_all()

        aktualisierte_zuordnungen = (
            db.query(ObjektKontakt).filter(ObjektKontakt.id.in_([item.id for item in zuordnungen])).all()
        )
        freigaben = (
            db.query(ObjektKontaktFreigabe)
            .filter(ObjektKontaktFreigabe.objekt_kontakt_id.in_([item.id for item in zuordnungen]))
            .all()
        )
        assert len(aktualisierte_zuordnungen) == 2
        assert all(
            zuordnung.zentraler_kontakt.telefone[0].nummer_normalisiert == telefon_normalisiert(nummer_b)
            for zuordnung in aktualisierte_zuordnungen
        )
        assert len(freigaben) == 2
        assert all(freigabe.ziel_wert == telefon_normalisiert(nummer_a) for freigabe in freigaben)
    finally:
        db.close()


def test_import_markiert_doppelte_oder_fremde_ids_als_konflikt_oder_fehler(client):
    user = _setup_user("kontakte_import_konflikt", "kontakt_verwalter")
    db = SessionLocal()
    set_tenant_context(db, None)
    try:
        foreign_org = FireDept(slug="kontakte-import-fremd", name="Import Fremd", color="#123456", bos="FW")
        db.add(foreign_org)
        db.flush()
        foreign = Kontakt(org_id=foreign_org.id, anzeigename="Fremder Kontakt")
        db.add(foreign)
        db.commit()
        duplicate = preview_import(
            db,
            user.org_id,
            [{"id": "77", "anzeigename": "A"}, {"id": "77", "anzeigename": "B"}],
        )
        foreign_row = preview_import(db, user.org_id, [{"id": str(foreign.id), "anzeigename": "Fremd"}])
        assert [item["status"] for item in duplicate] == ["konflikt", "konflikt"]
        assert foreign_row[0]["status"] == "fehler"
    finally:
        db.close()
